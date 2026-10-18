import html
import os
import re


from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

TITLE = "Расцепитель минимального напряжения"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TOPIC_DIR = os.path.dirname(SCRIPT_DIR)

# Корень сайта — папка над content/, как бы глубоко ни лежала тема.
_topic_parts = TOPIC_DIR.split(os.sep)
if "content" not in _topic_parts:
    raise SystemExit(f"Тема должна лежать внутри content/: {TOPIC_DIR}")
SITE_ROOT = os.sep.join(_topic_parts[: _topic_parts.index("content")])

DOCX_NAME = "Лекция — Расцепитель минимального напряжения.docx"
HTML_NAME = "Лекция (расцепитель).html"
IMAGE_ROOT = os.path.join(TOPIC_DIR, "images")

OUTPUT_DOCX = os.path.join(SCRIPT_DIR, DOCX_NAME)
OUTPUT_HTML = os.path.join(TOPIC_DIR, HTML_NAME)

# Пути в HTML строятся от корня сайта: content/groups/.../images и .../docs/....
TOPIC_URL = os.path.relpath(TOPIC_DIR, SITE_ROOT).replace(os.sep, "/")
IMAGE_URL = TOPIC_URL + "/images"
DOCX_URL = TOPIC_URL + "/docs/" + DOCX_NAME

TEXT = [
    ("p", "Расцепитель минимального напряжения — это электромагнитный аппарат, который удерживает коммутационный аппарат (автоматический выключатель, контактор или пускатель) во включённом положении, пока напряжение на его катушке достаточно. При недопустимом снижении напряжения расцепитель отпускает якорь, и аппарат отключает нагрузку. Так защищают оборудование от работы на пониженном напряжении: двигатель при этом перегревается, а электронные устройства дают сбои."),

    ("h", "Как он выглядит"),
    ("p", "В низковольтных щитах расцепитель выполняют в виде отдельного модуля, который встраивают в аппарат или присоединяют к нему сбоку. Автоматический выключатель без такого модуля защищает только от токов перегрузки и короткого замыкания."),
    ("img", "cb-2pole-din.jpg",
     "Автоматический выключатель IEK BA47-29 C16",
     "Фото: Kae, Wikimedia Commons, CC BY-SA 3.0"),
    ("p", "Модуль расцепителя имеет две клеммы для цепи управления и механически связан с замком аппарата. На контакторах роль расцепителя часто выполняет сама втягивающая катушка: при потере напряжения она отпускает якорь с главными контактами."),
    ("img", "contactor-iek.jpg",
     "Контактор IEK КМИ-11610: клеммы катушки А1 и А2",
     "Фото: Kae, Wikimedia Commons, CC BY-SA 3.0"),

    ("h", "Устройство и принцип действия"),
    ("p", "Основные части расцепителя: катушка с магнитопроводом, подвижный якорь, возвратная пружина и толкатель (ударник). Катушку питают от той же цепи, напряжение которой контролируют, либо от отдельного источника: пока напряжение в норме, якорь притянут и удерживает замок, а при снижении напряжения ниже уставки пружина отпускает якорь, и аппарат размыкается. После восстановления напряжения аппарат сам не включается — его взводят вручную."),
    ("p", "Значения напряжения, при которых расцепитель обязан удерживать и обязан отпускать аппарат, заданы стандартом."),
    ("table", [
        ["Напряжение на катушке", "Поведение расцепителя"],
        ["0,85 Uн и выше", "аппарат должен надёжно удерживаться"],
        ["0,35…0,7 Uн", "аппарат должен отключиться"],
        ["0 Uн", "аппарат остаётся отключённым"],
    ]),
    ("p", "В расчётах удобно принимать уставку отпускания U<sub>отп</sub> = 0,5 Uн. Сопротивление катушки растёт при нагреве, поэтому уставку проверяют и при рабочей температуре."),

    ("h", "Применение"),
    ("p", "Расцепитель минимального напряжения защищает электродвигатели и насосы от работы при пониженном напряжении, исключает самозапуск оборудования после исчезновения и восстановления напряжения, отключает нагрузки при переходе на резервное или аварийное питание и блокирует оборудование, если пропало питание цепей управления."),

    ("h", "Расчёт параметров катушки"),
    ("p", "Номинальный ток и сопротивление катушки определяют по её паспортным данным — номинальному напряжению Uн и мощности Pн:"),
    ("f", "I<sub>н</sub> = P<sub>н</sub> / U<sub>н</sub>"),
    ("f", "R<sub>к</sub> = U<sub>н</sub><sup>2</sup> / P<sub>н</sub>"),
    ("p", "Например, для катушки 24 В мощностью 3 Вт: Iн = 0,125 А, Rк = 192 Ом. Сопротивление катушки при нагреве растёт, и напряжение отпускания также увеличивается:"),
    ("f", "R(θ) = R<sub>20</sub> · [1 + α · (θ − 20)]"),
    ("p", "где α = 0,004 1/°C для меди. При нагреве до 40 °C сопротивление составляет 207,4 Ом, а напряжение отпускания — 12,96 В, то есть на 8 % выше значения при 20 °C. Катушку питают по проводам, поэтому учитывают потерю напряжения в линии:"),
    ("f", "ΔU = 2 · ρ · l · I / S"),

    ("h", "Отличие от других расцепителей"),
    ("p", "Независимый (шунтовой) расцепитель отключает аппарат при подаче напряжения на катушку, расцепитель максимального тока — при перегрузке или коротком замыкании, тепловой — при длительной перегрузке. Расцепитель минимального напряжения действует наоборот: он отключает аппарат, когда напряжение на катушке падает ниже уставки."),

    ("h", "Краткие выводы"),
    ("p", "Расцепитель минимального напряжения удерживает аппарат, пока есть напряжение, и отключает его при снижении напряжения ниже уставки: удержание — при 0,85 Uн, отпускание — при 0,35…0,7 Uн. После восстановления напряжения аппарат не включается сам — это защита от самозапуска. Расчёт сводится к определению тока и сопротивления катушки, проверке уставки при нагреве и проверке потери напряжения в линии."),
]


def add_run(paragraph, text, size=11, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def add_rich(paragraph, text, size=11, bold=False, italic=False, color=None):
    position = 0
    for match in re.finditer(r"(<sub>[^<]*</sub>|<sup>[^<]*</sup>)", text):
        if match.start() > position:
            add_run(paragraph, text[position:match.start()], size, bold, italic, color)
        token = match.group(0)
        run = add_run(paragraph, token[5:-6], size, bold, italic, color)
        if token.startswith("<sub"):
            run.font.subscript = True
        else:
            run.font.superscript = True
        position = match.end()
    if position < len(text):
        add_run(paragraph, text[position:], size, bold, italic, color)


def add_caption(document, text, credit):
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_together = True
    add_run(caption, text + " ", size=9, italic=True, color=RGBColor(90, 90, 90))
    add_run(caption, credit, size=8, color=RGBColor(120, 120, 120))


def make_docx(path):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    heading = document.styles["Heading 1"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(7)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.line_spacing = 1.15
    heading.paragraph_format.keep_with_next = True

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    add_run(title, TITLE, size=15, bold=True)

    caption_index = 0
    for kind, *value in TEXT:
        if kind == "h":
            document.add_paragraph(value[0], style="Heading 1")
        elif kind == "p":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1.0)
            add_rich(paragraph, value[0])
        elif kind == "ul":
            for item in value[0]:
                paragraph = document.add_paragraph(style="List Bullet")
                paragraph.paragraph_format.space_after = Pt(2)
                add_rich(paragraph, item)
        elif kind == "f":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True
            add_rich(paragraph, value[0])
        elif kind == "img":
            filename, alt, credit = value
            caption_index += 1
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.space_before = Pt(4)
            picture.paragraph_format.space_after = Pt(1)
            picture.paragraph_format.keep_with_next = True
            picture.add_run().add_picture(os.path.join(IMAGE_ROOT, filename), width=Cm(5))
            add_caption(document, f"Рисунок {caption_index} — {alt}.", credit)
        elif kind == "table":
            rows = value[0]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row in enumerate(rows):
                for col_index, cell_text in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(1)
                    add_run(paragraph, cell_text, size=9.5, bold=row_index == 0)
            spacer = document.add_paragraph()
            spacer.paragraph_format.space_after = Pt(0)

    document.core_properties.title = TITLE
    document.core_properties.subject = "Лекция"
    document.core_properties.author = "Электробаза"
    document.save(path)


def make_html(path):
    parts = []
    open_section = False
    caption_index = 0

    def close_section():
        nonlocal open_section
        if open_section:
            parts.append("</section>\n")
            open_section = False

    for kind, *value in TEXT:
        if kind == "h":
            close_section()
            parts.append("<section>\n")
            parts.append(f"  <h2>{html.escape(value[0], quote=False)}</h2>\n")
            open_section = True
        elif kind == "p":
            parts.append(f"  <p>{html.escape(value[0], quote=False)}</p>\n")
        elif kind == "ul":
            parts.append("  <ul>\n")
            for item in value[0]:
                parts.append(f"    <li>{html.escape(item, quote=False)}</li>\n")
            parts.append("  </ul>\n")
        elif kind == "f":
            parts.append(f'  <ul class="formulas">\n    <li>{html.escape(value[0], quote=False)}</li>\n  </ul>\n')
        elif kind == "img":
            filename, alt, credit = value
            caption_index += 1
            parts.append(
                "  <figure>\n"
                f'    <img src="{IMAGE_URL}/{filename}" alt="{html.escape(alt)}" />\n'
                f"    <figcaption>Рисунок {caption_index} — {html.escape(alt, quote=False)}. {html.escape(credit, quote=False)}</figcaption>\n"
                "  </figure>\n"
            )
        elif kind == "table":
            rows = value[0]
            parts.append("  <table>\n    <thead>\n      <tr>")
            for head in rows[0]:
                parts.append(f"<th>{html.escape(head, quote=False)}</th>")
            parts.append("</tr>\n    </thead>\n    <tbody>\n")
            for row in rows[1:]:
                parts.append("      <tr>")
                for cell in row:
                    parts.append(f"<td>{html.escape(cell, quote=False)}</td>")
                parts.append("</tr>\n")
            parts.append("    </tbody>\n  </table>\n")

    close_section()
    parts.append(
        '<section>\n  <h2>Источники иллюстраций</h2>\n'
        '  <ul>\n'
        '    <li><a href="https://commons.wikimedia.org/wiki/File:Circuit_breaker_2_pole_on_DIN_rail.JPG">Автоматический выключатель IEK BA47-29 C16</a> — Kae, CC BY-SA 3.0</li>\n'
        '    <li><a href="https://commons.wikimedia.org/wiki/File:Contactor_DIN_IEK.jpg">Контактор IEK КМИ-11610</a> — Kae, CC BY-SA 3.0</li>\n'
        '  </ul>\n'
        '  <p>Изображения получены из Wikimedia Commons и используются по указанным лицензиям.</p>\n'
        f'  <p><a href="{DOCX_URL}">Скачать лекцию в Word (.docx)</a></p>\n'
        '</section>\n'
    )
    with open(path, "w", encoding="utf-8", newline="\n") as file:
        file.write("".join(parts))


def main():
    os.makedirs(SCRIPT_DIR, exist_ok=True)
    os.makedirs(IMAGE_ROOT, exist_ok=True)
    for kind, *value in TEXT:
        if kind == "img":
            target = os.path.join(IMAGE_ROOT, value[0])
            if not os.path.exists(target):
                raise SystemExit(f"Нет файла изображения: {target}")
    make_docx(OUTPUT_DOCX)
    make_html(OUTPUT_HTML)
    words = 0
    for kind, *value in TEXT:
        if kind in {"p", "f"}:
            words += len(re.findall(r"[\wА-Яа-яЁё]+", value[0]))
        elif kind == "ul":
            words += sum(len(re.findall(r"[\wА-Яа-яЁё]+", item)) for item in value[0])
    print(f"DOCX: {OUTPUT_DOCX}")
    print(f"HTML: {OUTPUT_HTML}")
    print(f"Words: {words}; HTML bytes: {os.path.getsize(OUTPUT_HTML)}")


if __name__ == "__main__":
    main()
