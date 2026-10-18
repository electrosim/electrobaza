import html
import os

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt, RGBColor

TITLE = "Конституционный строй: форма правления, форма государственного устройства и политический режим"

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
TOPIC_DIR = os.path.dirname(SCRIPT_DIR)
_topic_parts = TOPIC_DIR.split(os.sep)
if "content" not in _topic_parts:
    raise SystemExit(f"Тема должна лежать внутри content/: {TOPIC_DIR}")
SITE_ROOT = os.sep.join(_topic_parts[: _topic_parts.index("content")])

TOPIC_URL = os.path.relpath(TOPIC_DIR, SITE_ROOT).replace(os.sep, "/")
IMAGE_URL = TOPIC_URL + "/images"
IMAGE_DIR = os.path.join(TOPIC_DIR, "images")
DOCX_NAME = "Лекция — Конституционный строй.docx"
HTML_NAME = "Лекция.html"
DOCX_URL = TOPIC_URL + "/docs/" + DOCX_NAME

OUTPUT_DOCX = os.path.join(SCRIPT_DIR, DOCX_NAME)
OUTPUT_HTML = os.path.join(TOPIC_DIR, HTML_NAME)

SCHEME_SVG = "schema-forma-gosudarstva-mobile.svg"
SCHEME_PNG = "schema-forma-gosudarstva.png"

TEXT = [
    ("p", "Конституционный строй — это система базовых принципов организации государства и общества, закреплённых Конституцией РФ 1993 г. (с поправками 2020 г.): народовластие, права и свободы человека, разделение властей, федерализм, правовое и социальное государство, идеологическое и политическое многообразие. Основы конституционного строя закреплены в главе 1 Конституции и защищены особым порядком пересмотра."),

    ("h", "Форма государства"),
    ("p", "Форма государства описывает, как организована и осуществляется государственная власть. Она включает три взаимосвязанные характеристики; их единство и образует форму конкретного государства."),
    ("img", SCHEME_SVG,
     "Три характеристики формы государства: форма правления, форма государственного устройства и политический режим"),

    ("h", "1. Форма правления"),
    ("p", "Форма правления — способ организации верховной государственной власти, порядок формирования её органов и их взаимодействия. Монархия передаёт власть по наследству, республика избирает главу государства на срок. Конституция РФ устанавливает республиканскую форму правления (ст. 1). По соотношению полномочий Президента, Правительства и Государственной Думы Россия является смешанной (полупрезидентской) республикой: Президент — глава государства и гарант Конституции (ст. 80), Правительство осуществляет исполнительную власть (ст. 110), Дума участвует в формировании Правительства и может выразить ему недоверие, а Президент вправе отправить Правительство в отставку."),
    ("table", [
        ["Форма правления", "Правительство и ответственность", "Примеры"],
        ["Президентская республика", "Формирует президент; ответственно перед президентом", "США, Бразилия"],
        ["Парламентская республика", "Формирует парламент; ответственно перед парламентом", "Германия, Италия"],
        ["Смешанная республика", "Формируют президент и парламент; ответственно перед обоими", "Россия, Франция"],
    ], "Таблица 1 — Формы правления"),

    ("h", "2. Форма государственного устройства"),
    ("p", "Форма государственного устройства — территориальная организация государства и соотношение полномочий центра и регионов. В унитарном государстве действует единая система власти, а регионы не обладают суверенитетом. Федерация состоит из субъектов, имеющих собственные органы и предметы ведения. Конфедерация — союз государств для ограниченных общих целей. Россия — федерация (ст. 1, 5): государственная целостность, единство системы власти, разграничение предметов ведения, равноправие субъектов. Субъекты перечислены в ст. 65, разграничение предметов ведения — ст. 71–73."),
    ("table", [
        ["Форма устройства", "Признаки", "Примеры"],
        ["Унитарное государство", "Единая конституция и система органов; регионы без суверенитета", "Франция, Япония"],
        ["Федерация", "Субъекты имеют органы власти и предметы ведения", "Россия, США, Германия"],
        ["Конфедерация", "Союз государств для ограниченных целей", "Исторические союзы; встречается редко"],
    ], "Таблица 2 — Формы государственного устройства"),

    ("h", "3. Политический режим"),
    ("p", "Политический режим — совокупность методов и средств осуществления государственной власти, положение личности, отношение к оппозиции. Конституция РФ закрепляет основы демократического режима: народ — носитель суверенитета и единственный источник власти (ст. 3); признаются идеологическое и политическое многообразие, многопартийность (ст. 13); гарантируются права и свободы, разделение властей (ст. 10), местное самоуправление (ст. 12); граждане участвуют в управлении через выборы и референдум."),
    ("p", "Недемократические режимы различаются степенью контроля. При авторитарном власть сосредоточена у одного лица или группы, оппозиция ограничена, но частная жизнь относительно свободна. При тоталитарном государство контролирует все сферы, существует единая обязательная идеология, подавляется инакомыслие."),

    ("h", "4. Как связаны три характеристики"),
    ("p", "Три характеристики взаимосвязаны, но не совпадают. Республика может быть демократической или авторитарной, федерация — демократической или централизованной, унитарное государство — демократическим. Поэтому о режиме нельзя судить только по форме правления: конституционная модель задаёт рамки, а реализация зависит от политической практики."),

    ("h", "Вывод"),
    ("p", "Россия по Конституции — демократическое федеративное правовое государство с республиканской формой правления. Это означает смешанную республику, федеративное устройство и демократический политический режим как конституционную основу. Юристу важно различать формальные признаки формы государства и их фактическую реализацию."),

    ("h", "Контрольные вопросы"),
    ("ol", [
        "Что такое конституционный строй и какие принципы относятся к его основам?",
        "Чем президентская республика отличается от парламентской?",
        "Почему Россию относят к смешанным республикам? Назовите не менее трёх признаков.",
        "Какие признаки отличают федерацию от унитарного государства?",
        "Может ли федерация быть недемократической? Приведите пример и объясните связь формы устройства и политического режима.",
    ]),
]


def add_run(paragraph, text, size=11, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def add_caption(document, text):
    caption = document.add_paragraph()
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_together = True
    add_run(caption, text, size=9, italic=True, color=RGBColor(90, 90, 90))


def make_docx(path):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.08
    normal.paragraph_format.space_after = Pt(3)

    heading = document.styles["Heading 1"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(6)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.line_spacing = 1.08
    heading.paragraph_format.keep_with_next = True

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(6)
    add_run(title, TITLE, size=14, bold=True)

    caption_index = 0
    for kind, *value in TEXT:
        if kind == "h":
            document.add_paragraph(value[0], style="Heading 1")
        elif kind == "p":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(0.9)
            add_run(paragraph, value[0])
        elif kind == "img":
            filename, alt = value[:2]
            caption_index += 1
            picture = document.add_paragraph()
            picture.alignment = WD_ALIGN_PARAGRAPH.CENTER
            picture.paragraph_format.space_before = Pt(4)
            picture.paragraph_format.space_after = Pt(1)
            picture.paragraph_format.keep_with_next = True
            picture.add_run().add_picture(os.path.join(IMAGE_DIR, SCHEME_PNG), width=Cm(12.5))
            add_caption(document, f"Рисунок {caption_index} — {alt}.")
        elif kind == "table":
            rows, caption_text = value
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            table.style = "Table Grid"
            table.autofit = True
            for row_index, row in enumerate(rows):
                for col_index, cell_text in enumerate(row):
                    cell = table.cell(row_index, col_index)
                    cell.text = ""
                    paragraph = cell.paragraphs[0]
                    paragraph.paragraph_format.space_after = Pt(1)
                    add_run(paragraph, cell_text, size=9.5, bold=row_index == 0)
            add_caption(document, caption_text)
        elif kind == "ol":
            for index, item in enumerate(value[0], 1):
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.left_indent = Cm(0.7)
                paragraph.paragraph_format.first_line_indent = Cm(-0.7)
                add_run(paragraph, f"{index}. {item}")

    document.core_properties.title = TITLE
    document.core_properties.subject = "Лекция"
    document.core_properties.author = "Электробаза"
    document.save(path)


def make_html(path):
    parts = []
    open_section = False
    caption_index = 0

    def close_section():
        nonlocal open_section
        if open_section:
            parts.append("</section>\n")
            open_section = False

    for kind, *value in TEXT:
        if kind == "h":
            close_section()
            parts.append("<section>\n")
            parts.append(f"  <h2>{html.escape(value[0], quote=False)}</h2>\n")
            open_section = True
        elif kind == "p":
            parts.append(f"  <p>{html.escape(value[0], quote=False)}</p>\n")
        elif kind == "img":
            filename, alt = value[:2]
            caption_index += 1
            parts.append(
                "  <figure>\n"
                f'    <img src="{IMAGE_URL}/{filename}" alt="{html.escape(alt)}" />\n'
                f"    <figcaption>Рисунок {caption_index} — {html.escape(alt, quote=False)}.</figcaption>\n"
                "  </figure>\n"
            )
        elif kind == "table":
            rows, caption_text = value
            parts.append(f"  <table>\n    <caption>{html.escape(caption_text, quote=False)}</caption>\n    <thead>\n      <tr>")
            for head in rows[0]:
                parts.append(f"<th>{html.escape(head, quote=False)}</th>")
            parts.append("</tr>\n    </thead>\n    <tbody>\n")
            for row in rows[1:]:
                parts.append("      <tr>")
                for cell in row:
                    parts.append(f"<td>{html.escape(cell, quote=False)}</td>")
                parts.append("</tr>\n")
            parts.append("    </tbody>\n  </table>\n")
        elif kind == "ol":
            parts.append("  <ol>\n")
            for item in value[0]:
                parts.append(f"    <li>{html.escape(item, quote=False)}</li>\n")
            parts.append("  </ol>\n")

    close_section()
    parts.append(
        f'<p><a href="{DOCX_URL}">Скачать лекцию в Word (.docx)</a></p>\n'
    )
    with open(path, "w", encoding="utf-8", newline="\n") as file:
        file.write("".join(parts))


def main():
    if not os.path.exists(os.path.join(IMAGE_DIR, SCHEME_PNG)):
        raise SystemExit(f"Нет схемы для Word: {os.path.join(IMAGE_DIR, SCHEME_PNG)}")
    make_docx(OUTPUT_DOCX)
    make_html(OUTPUT_HTML)
    print(f"DOCX: {OUTPUT_DOCX}")
    print(f"HTML: {OUTPUT_HTML}")


if __name__ == "__main__":
    main()
