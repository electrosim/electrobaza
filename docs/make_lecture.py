import html
import os
import re

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

TITLE = "Расчёт схемы с расцепителем минимального напряжения и датчиком освещённости"
SUBTITLE = "Краткая лекция к практической работе · МДК 04.02, 3-й семестр"

BLOCKS = [
    ("p", "Схема автоматического управления освещением включает датчик освещённости, который включает лампу при наступлении темноты, и расцепитель минимального напряжения, отключающий нагрузку при недопустимом снижении напряжения питания. Расчёт сводится к определению параметров катушки расцепителя, делителя с фоторезистором и сечения осветительной линии."),

    ("h", "1. Расцепитель минимального напряжения"),
    ("p", "Расцепитель удерживает аппарат во включённом положении, пока напряжение на катушке достаточно для притяжения якоря. При снижении напряжения до 0,35…0,7 U_{н} расцепитель отпускает якорь и отключает цепь. При напряжении 0,85 U_{н} и выше он должен надёжно удерживаться; в расчёте принимают уставку 0,5 U_{н}."),
    ("f", "I_{н} = P_{н} / U_{н}"),
    ("f", "R_{к} = U_{н}^{2} / P_{н}"),
    ("f", "U_{отп} = 0,5 · U_{н}"),
    ("f", "U_{вкл} = 0,85 · U_{н}"),
    ("p", "Для U_{н} = 24 В и P_{н} = 3 Вт: I_{н} = 0,125 А; R_{к} = 192 Ом; U_{отп} = 12 В; U_{вкл} = 20,4 В. Сопротивление медной катушки растёт при нагреве, поэтому уставку проверяют при рабочей температуре."),
    ("f", "R(θ) = R_{20} · [1 + α · (θ − 20)]"),
    ("f", "U_{отп}(θ) = I_{отп} · R(θ)"),
    ("p", "При 40 °C сопротивление катушки 207,4 Ом, а U_{отп} = 12,96 В — на 8 % выше, чем при 20 °C. Катушку питают по линии; потерю напряжения рассчитывают по формуле"),
    ("f", "ΔU = 2 · ρ · l · I / S"),
    ("p", "При l = 50 м, S = 1,5 мм² и I = 0,125 А потеря составляет 0,146 В, напряжение на катушке — 23,85 В, что выше 20,4 В. Суммарная мощность схемы складывается из мощности расцепителя, датчика и реле:"),
    ("f", "P_{Σ} = P_{н} + P_{д} + P_{рел}"),
    ("p", "Для примера P_{Σ} = 4,7 Вт, I_{Σ} = 0,196 А; подходит блок 24 В на 1,25 А."),

    ("h", "2. Датчик освещённости"),
    ("p", "Сопротивление фоторезистора R_{ф} уменьшается с ростом освещённости E. В делителе фоторезистор включают нижним плечом, а резистор R_{1} — верхним. Выходное напряжение снимают с фоторезистора:"),
    ("f", "U_{вых} = U_{пит} · R_{ф} / (R_{1} + R_{ф})"),
    ("p", "В темноте R_{ф} велико, поэтому U_{вых} близко к U_{пит}; при освещении напряжение уменьшается. Резистор R_{1} выбирают так, чтобы при требуемой освещённости включения выполнялось U_{вых} = U_{пор}:"),
    ("f", "R_{1} = R_{ф}(E_{вкл}) · (U_{пит} / U_{пор} − 1)"),
    ("p", "Для E_{вкл} = 150 лк R_{ф} = 5,5 кОм, U_{пит} = 12 В и U_{пор} = 2,5 В: R_{1} = 5,5 · 3,8 = 20,9 кОм, принимают 20 кОм. Фактический порог 2,5 В достигается при R_{ф} = 5,26 кОм, что соответствует E_{вкл} ≈ 164 лк."),
    ("p", "Чтобы лампа не мигала около порога, датчик имеет гистерезис h:"),
    ("f", "U_{отк} = (1 − h) · U_{вкл}"),
    ("p", "При h = 25 % U_{отк} = 1,875 В; этому соответствуют R_{ф} = 3,70 кОм и E_{отк} ≈ 258 лк. Гистерезис по освещённости — 94 лк."),

    ("h", "3. Осветительная линия"),
    ("p", "Ток осветительной линии и потерю напряжения определяют по формулам:"),
    ("f", "I_{л} = P_{л} / (U_{с} · cos φ)"),
    ("f", "ΔU = 2 · ρ · L · I_{л} / S"),
    ("f", "ΔU, % = 100 · ΔU / U_{с}"),
    ("p", "Для активной нагрузки cos φ = 1. При P_{л} = 3,0 кВт I_{л} = 13,04 А. При L = 60 м: S = 1,5 мм² — 7,94 %; 2,5 — 4,76 %; 4 — 2,98 %; 6 — 1,98 %."),
    ("p", "Условию ΔU_{доп} = 3 % удовлетворяют 4 и 6 мм². Принимают 4 мм², а при возможности — 6 мм² с запасом. Автоматический выключатель B16 согласован с кабелем 4 мм². Расчёт выполняют при 20 °C; при нагреве сопротивление меди растёт, поэтому реальная потеря больше, и запас по сечению полезен."),

    ("h", "4. Порядок расчёта"),
    ("ol", [
        "Выбрать данные варианта.",
        "Рассчитать цепь расцепителя и проверить удержание при нагреве.",
        "Рассчитать R_{1} и пороги датчика.",
        "Выбрать сечение линии по 3 % и проверить по допустимому току; сделать вывод.",
    ]),

    ("h", "Вывод"),
    ("p", "Схема работоспособна, если расцепитель удерживается при 0,85 U_{н}, пороги датчика соответствуют заданной освещённости, а потеря напряжения не превышает 3 %. Если условие не выполняется, увеличивают сечение кабеля или корректируют R_{1}."),
]

TOKEN = re.compile(r"(_\{[^}]*\}|\^\{[^}]*\})")


def add_run(paragraph, text, size=11, bold=False, italic=False, color=None):
    run = paragraph.add_run(text)
    run.font.name = "Times New Roman"
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    if color is not None:
        run.font.color.rgb = color
    return run


def add_rich(paragraph, text, size=11, bold=False, italic=False, color=None):
    position = 0
    for match in TOKEN.finditer(text):
        if match.start() > position:
            add_run(paragraph, text[position:match.start()], size, bold, italic, color)
        token = match.group(0)
        content = token[2:-1]
        run = add_run(paragraph, content, size, bold, italic, color)
        if token.startswith("_{"):
            run.font.subscript = True
        else:
            run.font.superscript = True
        position = match.end()
    if position < len(text):
        add_run(paragraph, text[position:], size, bold, italic, color)


def make_docx(path):
    document = Document()
    section = document.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2)
    section.right_margin = Cm(2)

    normal = document.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(11)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(3)

    heading = document.styles["Heading 1"]
    heading.font.name = "Times New Roman"
    heading.font.size = Pt(11.5)
    heading.font.bold = True
    heading.font.color.rgb = RGBColor(0, 0, 0)
    heading.paragraph_format.space_before = Pt(8)
    heading.paragraph_format.space_after = Pt(2)
    heading.paragraph_format.line_spacing = 1.15
    heading.paragraph_format.keep_with_next = True

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(2)
    add_run(title, TITLE, size=15, bold=True)

    subtitle = document.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle.paragraph_format.space_after = Pt(8)
    add_run(subtitle, SUBTITLE, size=10, italic=True, color=RGBColor(90, 90, 90))
    ppr = subtitle._p.get_or_add_pPr()
    border = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "999999")
    border.append(bottom)
    ppr.append(border)

    for kind, value in BLOCKS:
        if kind == "h":
            paragraph = document.add_paragraph(value, style="Heading 1")
        elif kind == "p":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            paragraph.paragraph_format.first_line_indent = Cm(1.0)
            add_rich(paragraph, value)
        elif kind == "f":
            paragraph = document.add_paragraph()
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            paragraph.paragraph_format.first_line_indent = Cm(0)
            paragraph.paragraph_format.space_before = Pt(2)
            paragraph.paragraph_format.space_after = Pt(2)
            paragraph.paragraph_format.keep_together = True
            add_rich(paragraph, value)
        elif kind == "ol":
            for index, item in enumerate(value, 1):
                paragraph = document.add_paragraph()
                paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
                paragraph.paragraph_format.left_indent = Cm(0.75)
                paragraph.paragraph_format.first_line_indent = Cm(-0.75)
                add_rich(paragraph, f"{index}. {item}")

    document.core_properties.title = TITLE
    document.core_properties.subject = "Краткая лекция к практической работе"
    document.core_properties.author = "Электробаза"
    document.save(path)


def notation_to_html(text):
    escaped = html.escape(text, quote=False)
    escaped = re.sub(r"_\{([^}]*)\}", r"<sub>\1</sub>", escaped)
    escaped = re.sub(r"\^\{([^}]*)\}", r"<sup>\1</sup>", escaped)
    return escaped


def plain_length(text):
    return len(TOKEN.sub(lambda m: m.group(0)[2:-1], text))


def make_html(path):
    parts = []
    open_section = False
    in_formula_list = False

    def close_formulas():
        nonlocal in_formula_list
        if in_formula_list:
            parts.append("  </ul>\n")
            in_formula_list = False

    for kind, value in BLOCKS:
        if kind != "f":
            close_formulas()
        if kind == "h":
            if open_section:
                parts.append("</section>\n")
            parts.append("<section>\n")
            parts.append(f"  <h2>{notation_to_html(value)}</h2>\n")
            open_section = True
        elif kind == "p":
            parts.append(f"  <p>{notation_to_html(value)}</p>\n")
        elif kind == "f":
            if not in_formula_list:
                parts.append('  <ul class="formulas">\n')
                in_formula_list = True
            cls = ' class="small"' if plain_length(value) >= 28 else ""
            parts.append(f"    <li{cls}>{notation_to_html(value)}</li>\n")
        elif kind == "ol":
            parts.append("  <ol>\n")
            for item in value:
                parts.append(f"    <li>{notation_to_html(item)}</li>\n")
            parts.append("  </ol>\n")

    close_formulas()
    if open_section:
        parts.append("</section>\n")

    parts.append('<section>\n  <h2>Материалы</h2>\n  <p>Подробные расчёты по этой теме — в <a href="#/material/1eo-24/mdk-04-02-1eo-24-s3/undervoltage-release-light-sensor/undervoltage-release-light-sensor-practice">практической работе</a>.</p>\n  <p><a href="docs/lecture-undervoltage-release-light-sensor.docx">Скачать лекцию в Word (.docx)</a></p>\n</section>\n')
    with open(path, "w", encoding="utf-8", newline="\n") as file:
        file.write("".join(parts))


def main():
    os.makedirs("docs", exist_ok=True)
    docx_path = os.path.join("docs", "lecture-undervoltage-release-light-sensor.docx")
    html_path = os.path.join(
        "content", "groups", "1eo-24", "mdk-04-02", "semester-3",
        "undervoltage-release-light-sensor", "lecture.html"
    )
    make_docx(docx_path)
    make_html(html_path)

    text = " ".join(
        value if isinstance(value, str) else " ".join(value)
        for kind, value in BLOCKS
        if kind != "h"
    )
    words = len(re.findall(r"[\wА-Яа-яЁё]+", TOKEN.sub(lambda m: m.group(0)[2:-1], text)))
    print(f"DOCX: {docx_path}")
    print(f"HTML: {html_path}")
    print(f"Words: {words}; HTML chars: {len(open(html_path, encoding='utf-8').read())}")


if __name__ == "__main__":
    main()
